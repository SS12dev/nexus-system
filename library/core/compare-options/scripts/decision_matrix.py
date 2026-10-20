"""Weighted decision matrix with a sensitivity check. See SKILL.md for the JSON.

python decision_matrix.py options.json [--md out/comparison.md] [--xlsx out/comparison.xlsx]
"""

import argparse
import json
from pathlib import Path


def normalise(criteria: list[dict], options: list[dict]) -> dict[str, dict[str, float]]:
    out: dict[str, dict[str, float]] = {o["name"]: {} for o in options}
    for c in criteria:
        name = c["name"]
        vals = [float(o["values"][name]) for o in options if o["values"].get(name) is not None]
        lo, hi = c.get("scale") or (min(vals, default=0), max(vals, default=0))
        for o in options:
            v = o["values"].get(name)
            if v is None:
                out[o["name"]][name] = 0.0  # missing data scores as worst; say so in the report
                continue
            x = 0.5 if hi == lo else (float(v) - lo) / (hi - lo)
            out[o["name"]][name] = 1 - x if c.get("better", "higher") == "lower" else x
    return out


def totals(criteria: list[dict], norm: dict, weights: dict[str, float]) -> dict[str, float]:
    wsum = sum(weights.values()) or 1
    return {
        o: sum(norm[o][c["name"]] * weights[c["name"]] for c in criteria) / wsum * 100 for o in norm
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("--md")
    ap.add_argument("--xlsx")
    a = ap.parse_args()
    spec = json.loads(Path(a.spec).read_text("utf-8"))
    criteria, options = spec["criteria"], spec["options"]
    for o in options:
        unknown = set(o.get("values", {})) - {c["name"] for c in criteria}
        if unknown:
            raise SystemExit(f"{o['name']}: values for unknown criteria {sorted(unknown)}")
    ranked = [o for o in options if not o.get("must_fail")]
    out_of_race = [o for o in options if o.get("must_fail")]
    if not ranked:
        raise SystemExit("every option fails a must-have")
    weights = {c["name"]: float(c["weight"]) for c in criteria}
    norm = normalise(criteria, ranked)
    score = totals(criteria, norm, weights)
    order = sorted(score, key=score.get, reverse=True)
    winner = order[0]

    lines = [
        "| Rank | Option | Score | "
        + " | ".join(f"{c['name']} ({c['weight']:g})" for c in criteria)
        + " |"
    ]
    lines.append("| --- | --- | --- |" + " --- |" * len(criteria))
    by_name = {o["name"]: o for o in options}
    for i, name in enumerate(order, 1):
        cells = [str(by_name[name]["values"].get(c["name"], "–")) for c in criteria]
        lines.append(f"| {i} | {name} | {score[name]:.0f} | " + " | ".join(cells) + " |")
    for o in out_of_race:
        cells = [str(o["values"].get(c["name"], "–")) for c in criteria]
        lines.append(
            f"| – | {o['name']} (fails: {o['must_fail']}) | – | " + " | ".join(cells) + " |"
        )
    report = ["\n".join(lines), ""]
    for name in order:
        best = max(criteria, key=lambda c: norm[name][c["name"]])["name"]
        worst = min(criteria, key=lambda c: norm[name][c["name"]])["name"]
        report.append(f"- {name}: strongest on {best}, weakest on {worst}")

    report += ["", "Sensitivity (each weight moved by 20 points, others rescaled):"]
    stable = True
    for c in criteria:
        for delta in (-20, 20):
            w = dict(weights)
            w[c["name"]] = max(0.0, w[c["name"]] + delta)
            s = totals(criteria, norm, w)
            top = max(s, key=s.get)
            if top != winner:
                stable = False
                report.append(f"- {c['name']} {delta:+d}: {top} wins instead of {winner}")
    if stable:
        report.append(f"- {winner} stays first under every change: a robust choice.")
    missing = [
        (o["name"], c["name"])
        for o in ranked
        for c in criteria
        if o["values"].get(c["name"]) is None
    ]
    if missing:
        report.append(
            "\nMissing values (scored as worst): " + ", ".join(f"{o}/{c}" for o, c in missing)
        )
    text = "\n".join(report) + "\n"
    print(text)
    if a.md:
        Path(a.md).parent.mkdir(parents=True, exist_ok=True)
        Path(a.md).write_text(text, "utf-8")
    if a.xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Comparison"
        ws.append(["Option", "Score"] + [c["name"] for c in criteria] + ["Notes"])
        ws.append(["Weight", None] + [c["weight"] for c in criteria] + [None])
        for name in order:
            o = by_name[name]
            ws.append(
                [name, round(score[name], 1)]
                + [o["values"].get(c["name"]) for c in criteria]
                + [o.get("notes")]
            )
        for o in out_of_race:
            ws.append(
                [o["name"], None]
                + [o["values"].get(c["name"]) for c in criteria]
                + [f"fails: {o['must_fail']}"]
            )
        for cell in ws[1] + ws[2]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                12, min(40, max(len(str(c.value or "")) for c in col) + 2)
            )
        ws.freeze_panes = "B3"
        Path(a.xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(a.xlsx)
        print(f"Saved {a.xlsx}")


if __name__ == "__main__":
    main()
