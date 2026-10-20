"""Convert tables and documents between formats; the format comes from the extensions.

python convert.py INPUT OUTPUT [--sheet NAME] [--limit 200]
Tables: .csv .tsv .xlsx .json .jsonl -> .csv .tsv .xlsx .json .jsonl .md
Documents: .docx .pptx .pdf -> .md / .txt
"""

import argparse
import csv
import json
from pathlib import Path

import pandas as pd

TABLES = {".csv", ".tsv", ".xlsx", ".xlsm", ".json", ".jsonl"}
DOCS = {".docx", ".pptx", ".pdf"}


def read_csv(path: str, tsv: bool) -> pd.DataFrame:
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sep = "\t" if tsv else None
    if sep is None:
        try:
            sep = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            sep = ","
    from io import StringIO

    return pd.read_csv(StringIO(text), sep=sep)


def read_tables(path: str, sheet: str | None) -> dict[str, pd.DataFrame]:
    s = Path(path).suffix.lower()
    if s in (".csv", ".tsv"):
        return {Path(path).stem: read_csv(path, s == ".tsv")}
    if s in (".xlsx", ".xlsm"):
        if sheet:
            return {sheet: pd.read_excel(path, sheet_name=sheet)}
        return pd.read_excel(path, sheet_name=None)
    if s == ".jsonl":
        return {Path(path).stem: pd.read_json(path, lines=True)}
    data = json.loads(Path(path).read_text("utf-8"))
    if isinstance(data, dict) and all(isinstance(v, list) for v in data.values()):
        return {k: pd.json_normalize(v) for k, v in data.items()}
    return {Path(path).stem: pd.json_normalize(data)}


def md_table(df: pd.DataFrame, limit: int) -> str:
    def esc(v) -> str:
        return "" if pd.isna(v) else str(v).replace("|", "\\|").replace("\n", " ")

    head = "| " + " | ".join(esc(c) for c in df.columns) + " |"
    sep = "| " + " | ".join("---" for _ in df.columns) + " |"
    rows = [
        "| " + " | ".join(esc(v) for v in r) + " |" for r in df.head(limit).itertuples(index=False)
    ]
    more = [f"\n({len(df) - limit} more rows not shown)"] if len(df) > limit else []
    return "\n".join([head, sep, *rows, *more])


def write_tables(tables: dict[str, pd.DataFrame], out: str, limit: int) -> None:
    s = Path(out).suffix.lower()
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    first = next(iter(tables.values()))
    if s in (".csv", ".tsv"):
        if len(tables) > 1:
            print(
                f"Several sheets; writing the first "
                f"({next(iter(tables))}). Use --sheet for another."
            )
        first.to_csv(out, index=False, sep="\t" if s == ".tsv" else ",")
    elif s in (".xlsx", ".xlsm"):
        with pd.ExcelWriter(out, engine="openpyxl") as xw:
            for name, df in tables.items():
                df.to_excel(xw, sheet_name=str(name)[:31], index=False)
    elif s == ".json":
        payload = {
            k: json.loads(v.to_json(orient="records", date_format="iso")) for k, v in tables.items()
        }
        Path(out).write_text(
            json.dumps(
                payload if len(tables) > 1 else next(iter(payload.values())),
                indent=1,
                ensure_ascii=False,
            ),
            "utf-8",
        )
    elif s == ".jsonl":
        first.to_json(out, orient="records", lines=True, date_format="iso", force_ascii=False)
    elif s in (".md", ".txt"):
        parts = [
            (f"## {k}\n\n" if len(tables) > 1 else "") + md_table(v, limit)
            for k, v in tables.items()
        ]
        Path(out).write_text("\n\n".join(parts) + "\n", "utf-8")
    else:
        raise SystemExit(f"can't write tables as {s}")
    for name, df in tables.items():
        print(f"{name}: {len(df):,} rows x {df.shape[1]} columns")
    print(f"-> {out}")


def docx_to_md(path: str) -> str:
    from docx import Document

    doc = Document(path)
    out = []
    body = doc.element.body
    tables = {t._tbl: t for t in doc.tables}
    paras = {p._p: p for p in doc.paragraphs}
    for el in body.iterchildren():
        if el in paras:
            p = paras[el]
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                out.append("#" * min(int(level), 6) + " " + text)
            elif style == "title":
                out.append("# " + text)
            elif "list" in style:
                out.append(("1. " if "number" in style else "- ") + text)
            else:
                out.append(text)
        elif el in tables:
            rows = [[c.text.strip() for c in r.cells] for r in tables[el].rows]
            if rows:
                out.append(md_table(pd.DataFrame(rows[1:], columns=rows[0]), 10_000))
    return "\n\n".join(out)


def pptx_to_md(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = (
            slide.shapes.title.text.strip()
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame
            else ""
        )
        out.append(f"## Slide {i}" + (f": {title}" if title else ""))
        for sh in slide.shapes:
            if sh == slide.shapes.title:
                continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append("  " * p.level + "- " + t)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                out.append(md_table(pd.DataFrame(rows[1:], columns=rows[0]), 10_000))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        out.append("")
    return "\n".join(out)


def pdf_to_md(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = [
        f"## Page {i}\n\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1)
    ]
    text = "\n\n".join(pages)
    if not any((p.extract_text() or "").strip() for p in reader.pages):
        print("No text layer: this PDF is probably scanned images (no OCR here).")
    return text


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--sheet")
    ap.add_argument("--limit", type=int, default=200)
    a = ap.parse_args()
    src, dst = Path(a.input).suffix.lower(), Path(a.output).suffix.lower()
    if src in TABLES:
        write_tables(read_tables(a.input, a.sheet), a.output, a.limit)
    elif src in DOCS:
        if dst not in (".md", ".txt"):
            raise SystemExit(f"{src} converts to .md or .txt here; see SKILL.md for other targets")
        text = {".docx": docx_to_md, ".pptx": pptx_to_md, ".pdf": pdf_to_md}[src](a.input)
        Path(a.output).parent.mkdir(parents=True, exist_ok=True)
        Path(a.output).write_text(text + "\n", "utf-8")
        print(f"-> {a.output} ({len(text):,} characters)")
        print(text[:600])
    else:
        raise SystemExit(f"don't know how to convert {src}")


if __name__ == "__main__":
    main()
