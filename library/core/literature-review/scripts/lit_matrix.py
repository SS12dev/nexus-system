"""Literature matrix from a JSON list of paper notes. See SKILL.md for the fields.

python lit_matrix.py build papers.json out/matrix.xlsx [--md out/matrix.md]
python lit_matrix.py coverage papers.json --rows method --cols modality
python lit_matrix.py timeline papers.json --by method
"""

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

COLUMNS = [
    ("year", "Year", 7),
    ("title", "Title", 50),
    ("authors", "Authors", 22),
    ("venue", "Venue", 14),
    ("citations", "Citations", 10),
    ("problem", "Problem", 28),
    ("approach", "Approach", 28),
    ("data", "Data", 20),
    ("metrics", "Metrics", 18),
    ("findings", "Key findings", 50),
    ("limitations", "Limitations", 40),
    ("tags", "Tags", 30),
    ("url", "Link", 36),
]


def load(path: str) -> list[dict]:
    data = json.loads(Path(path).read_text("utf-8"))
    papers = data.get("papers", data) if isinstance(data, dict) else data
    if not isinstance(papers, list) or not all(isinstance(p, dict) for p in papers):
        raise SystemExit("papers.json must be a list of objects (or {'papers': [...]})")
    seen, out = set(), []
    for p in papers:
        key = " ".join(str(p.get("title", "")).lower().split())
        if not key:
            raise SystemExit(f"a paper has no title: {p}")
        if key in seen:
            continue  # the same paper found by two searches
        seen.add(key)
        out.append(p)
    return sorted(out, key=lambda p: (-int(p.get("year") or 0), str(p.get("title"))))


def tag(p: dict, name: str) -> list[str]:
    v = (p.get("tags") or {}).get(name, p.get(name))
    if v in (None, ""):
        return []
    return [str(x) for x in v] if isinstance(v, list) else [str(v)]


def cell(p: dict, key: str) -> str:
    v = p.get(key)
    if key == "tags" and isinstance(v, dict):
        return "; ".join(f"{k}: {', '.join(tag(p, k))}" for k in v)
    if isinstance(v, list):
        return ", ".join(map(str, v))
    return "" if v is None else str(v)


def build(papers: list[dict], out: str, md: str | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Papers"
    ws.append([h for _, h, _ in COLUMNS])
    for c, (_, _, width) in zip(ws[1], COLUMNS, strict=True):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.column_dimensions[c.column_letter].width = width
    for p in papers:
        ws.append([p.get(k) if k in ("year", "citations") else cell(p, k) for k, _, _ in COLUMNS])
        link = ws.cell(row=ws.max_row, column=len(COLUMNS))
        if link.value:
            link.hyperlink = link.value
            link.style = "Hyperlink"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{len(papers)} papers -> {out}")
    if md:
        keys = ["year", "title", "venue", "approach", "findings", "limitations"]
        lines = ["| Year | Paper | Venue | Approach | Key findings | Limitations |"]
        lines.append("| --- | --- | --- | --- | --- | --- |")
        for p in papers:
            row = [cell(p, k).replace("|", "/").replace("\n", " ") for k in keys]
            if p.get("url"):
                row[1] = f"[{row[1]}]({p['url']})"
            lines.append("| " + " | ".join(row) + " |")
        Path(md).parent.mkdir(parents=True, exist_ok=True)
        Path(md).write_text("\n".join(lines) + "\n", "utf-8")
        print(f"Markdown table -> {md}")


def coverage(papers: list[dict], rows: str, cols: str) -> None:
    grid: dict[tuple[str, str], int] = Counter()
    rvals, cvals = Counter(), Counter()
    for p in papers:
        for r in tag(p, rows) or ["(none)"]:
            for c in tag(p, cols) or ["(none)"]:
                grid[r, c] += 1
                rvals[r] += 1
                cvals[c] += 1
    rs, cs = [r for r, _ in rvals.most_common()], [c for c, _ in cvals.most_common()]
    print(f"| {rows} \\ {cols} | " + " | ".join(cs) + " |")
    print("| --- |" + " --- |" * len(cs))
    for r in rs:
        print(f"| {r} | " + " | ".join(str(grid[r, c] or "·") for c in cs) + " |")
    empty = [(r, c) for r in rs for c in cs if not grid[r, c]]
    thin = [(r, c) for r in rs for c in cs if grid[r, c] == 1]
    print(f"\nEmpty cells (candidate gaps, check you searched there): {len(empty)}")
    for r, c in empty:
        print(f"- {rows}={r} x {cols}={c}")
    if thin:
        print(f"Cells with a single paper: {len(thin)}")
        for r, c in thin:
            print(f"- {rows}={r} x {cols}={c}")


def timeline(papers: list[dict], by: str) -> None:
    counts: dict[str, Counter] = defaultdict(Counter)
    years = sorted({int(p["year"]) for p in papers if p.get("year")})
    for p in papers:
        if p.get("year"):
            for v in tag(p, by) or ["(none)"]:
                counts[v][int(p["year"])] += 1
    print(f"| {by} | " + " | ".join(map(str, years)) + " | total |")
    print("| --- |" + " --- |" * (len(years) + 1))
    for v, c in sorted(counts.items(), key=lambda kv: -sum(kv[1].values())):
        print(f"| {v} | " + " | ".join(str(c[y] or "·") for y in years) + f" | {sum(c.values())} |")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="papers.json -> .xlsx matrix (and Markdown)")
    b.add_argument("papers")
    b.add_argument("out")
    b.add_argument("--md")
    c = sub.add_parser("coverage", help="count papers for each pair of tag values")
    c.add_argument("papers")
    c.add_argument("--rows", required=True)
    c.add_argument("--cols", required=True)
    t = sub.add_parser("timeline", help="papers per year for each value of a tag")
    t.add_argument("papers")
    t.add_argument("--by", required=True)
    a = ap.parse_args()
    papers = load(a.papers)
    if a.cmd == "build":
        build(papers, a.out, a.md)
    elif a.cmd == "coverage":
        coverage(papers, a.rows, a.cols)
    else:
        timeline(papers, a.by)


if __name__ == "__main__":
    main()
