"""Build an .xlsx workbook from JSON with openpyxl. See SKILL.md for the format.

python build_workbook.py book.json out.xlsx
"""

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="EEF2F7")
HEADER_FONT = Font(bold=True, color="0F2A4A")
THIN = Side(style="thin", color="C8CED8")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ISO_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?$")
TOTALS = {"sum": "SUM", "average": "AVERAGE", "min": "MIN", "max": "MAX", "count": "COUNT"}


def cell_value(v):
    if isinstance(v, str):
        if ISO_DATE.match(v):
            return date.fromisoformat(v)
        if ISO_DATETIME.match(v):
            return datetime.fromisoformat(v)
    return v


def add_sheet(wb: Workbook, spec: dict) -> None:
    ws = wb.create_sheet(str(spec.get("name", f"Sheet{len(wb.sheetnames) + 1}"))[:31])
    columns = spec.get("columns")
    rows = spec.get("rows", [])
    if not columns:
        for row in rows:
            ws.append([cell_value(v) for v in row])
        return
    headers = [c["header"] for c in columns]
    ws.append(headers)
    for c_i in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=c_i)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for row in rows:
        ws.append([cell_value(v) for v in row])
    last = ws.max_row
    for c_i, col in enumerate(columns, 1):
        letter = get_column_letter(c_i)
        fmt = col.get("format")
        for r in range(2, last + 1):
            cell = ws.cell(row=r, column=c_i)
            if fmt:
                cell.number_format = fmt
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
        longest = max(
            [len(str(col["header"]))]
            + [
                len(str(ws.cell(row=r, column=c_i).value or ""))
                for r in range(2, min(last, 200) + 1)
            ]
        )
        ws.column_dimensions[letter].width = col.get("width") or min(max(10, longest + 2), 60)
    if spec.get("totals") and last >= 2:
        total_row = last + 1
        ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
        for header, fn in spec["totals"].items():
            if header not in headers:
                raise SystemExit(f"totals: no column {header!r}")
            c_i = headers.index(header) + 1
            letter = get_column_letter(c_i)
            cell = ws.cell(
                row=total_row, column=c_i, value=f"={TOTALS[fn]}({letter}2:{letter}{last})"
            )
            cell.font = Font(bold=True)
            cell.border = Border(top=THIN)
            if columns[c_i - 1].get("format"):
                cell.number_format = columns[c_i - 1]["format"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
    if chart := spec.get("chart"):
        add_chart(ws, chart, headers, last)


def add_chart(ws, spec: dict, headers: list[str], last: int) -> None:
    kind = spec.get("type", "col")
    chart = {"bar": BarChart, "col": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    if kind == "bar":
        chart.type = "bar"
    elif kind == "col":
        chart.type = "col"
    chart.title = spec.get("title")
    chart.height, chart.width = 9, 16
    cat_col = headers.index(spec["categories"]) + 1
    for name in spec.get("values", []):
        col = headers.index(name) + 1
        data = Reference(ws, min_col=col, min_row=1, max_row=last)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=cat_col, min_row=2, max_row=last))
    ws.add_chart(chart, f"{get_column_letter(len(headers) + 2)}2")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("spec", type=Path)
    ap.add_argument("output", type=Path)
    a = ap.parse_args()
    spec = json.loads(a.spec.read_text("utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in spec.get("sheets", []):
        add_sheet(wb, sheet)
    if not wb.sheetnames:
        raise SystemExit("spec has no sheets")
    a.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.output)
    print(f"Wrote {a.output} ({', '.join(wb.sheetnames)})")


if __name__ == "__main__":
    main()
