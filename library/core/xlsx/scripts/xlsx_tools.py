"""Inspect a workbook: sheets, sizes, headers, formulas, merged cells and charts.

python xlsx_tools.py inspect IN [--rows 5]
"""

import argparse
import json

from openpyxl import load_workbook


def inspect(a):
    wb = load_workbook(a.input)
    cached = load_workbook(a.input, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        formulas = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append(f"{c.coordinate}: {c.value}")
        preview = [
            [
                cached[ws.title].cell(row=r, column=c).value
                for c in range(1, min(ws.max_column, 12) + 1)
            ]
            for r in range(1, min(ws.max_row, a.rows) + 1)
        ]
        sheets.append(
            {
                "name": ws.title,
                "dimensions": ws.dimensions,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "preview": preview,
                "formulas": formulas[:30]
                + ([f"... {len(formulas) - 30} more"] if len(formulas) > 30 else []),
                "merged": [str(m) for m in ws.merged_cells.ranges][:20],
                "charts": len(ws._charts),
                "freeze_panes": ws.freeze_panes,
            }
        )
    print(json.dumps({"sheets": sheets}, indent=2, default=str))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("inspect")
    p.add_argument("input")
    p.add_argument("--rows", type=int, default=5)
    p.set_defaults(fn=inspect)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
