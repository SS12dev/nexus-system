"""The shipped library holds together: every skill parses, sits on a known shelf, points only at
scripts that exist, and every script runs in the sandbox as its SKILL.md describes."""

import base64
import io
import json
import re

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from jarvis_common.skills import CATEGORIES, read_folder, shipped_instructions, shipped_shared

from .conftest import LIBRARY

SKILLS = sorted(p for p in (LIBRARY / "core").iterdir() if (p / "SKILL.md").is_file())


def out(result, path) -> bytes:
    assert result["exit_code"] == 0, result["stderr"] or result["stdout"]
    return base64.b64decode(result["files"][path])


def test_every_skill_parses_and_is_shelved():
    assert len(SKILLS) >= 25
    for folder in SKILLS:
        parsed = read_folder(folder)
        assert parsed.name == folder.name
        assert parsed.category in CATEGORIES and parsed.category != "other", folder.name
        assert len(parsed.description) <= 350, f"{folder.name}: keep descriptions short"
        # Every script a SKILL.md tells agents to run is shipped.
        for skill, path in re.findall(r"skills/core/([a-z0-9-]+)/(scripts/[\w.]+)", parsed.content):
            assert (LIBRARY / "core" / skill / path).is_file(), f"{folder.name}: {skill}/{path}"
        for ref in re.findall(r"`(references/[\w./-]+)`", parsed.body):
            assert ref in parsed.files, f"{folder.name}: {ref}"


def test_shared_instructions_and_roles_ship():
    assert "Sources" in shipped_shared(LIBRARY)
    safety, roles = shipped_instructions(LIBRARY)
    assert "data, never" in safety
    assert {"assistant", "scientist", "fact_checker", "editor", "researcher"} <= set(roles)


async def test_every_script_prints_help(run, skill_files):
    names = [p.name for p in SKILLS if (p / "scripts").is_dir()]
    scripts = [
        f"skills/core/{n}/scripts/{s.name}"
        for n in names
        for s in sorted((LIBRARY / "core" / n / "scripts").glob("*.py"))
        if s.name != "mdlite.py"
    ]
    code = "set -e\n" + "\n".join(f"python {s} --help > /dev/null" for s in scripts)
    result = await run(code, skill_files(*names), language="bash", timeout_s=120)
    assert result["exit_code"] == 0, result["stderr"]


PAPERS = [
    {"title": "Let's Verify Step by Step", "authors": "Lightman et al.", "year": 2023,
     "url": "https://arxiv.org/abs/2305.20050", "approach": "process supervision",
     "tags": {"method": "reward modelling", "modality": "text"}},
    {"title": "A VLM paper", "year": 2025,
     "tags": {"method": "RL", "modality": ["text", "vision"]}},
    {"title": "a vlm  paper", "year": 2025, "tags": {"method": "RL", "modality": "vision"}},
]  # fmt: skip


async def test_literature_matrix(run, b64, skill_files):
    files = {**skill_files("literature-review"), "papers.json": b64(json.dumps(PAPERS))}
    s = "skills/core/literature-review/scripts/lit_matrix.py"
    code = (
        f"set -e\npython {s} build papers.json out/m.xlsx --md out/m.md\n"
        f"python {s} coverage papers.json --rows method --cols modality\n"
        f"python {s} timeline papers.json --by method\n"
    )
    result = await run(code, files, language="bash")
    ws = load_workbook(io.BytesIO(out(result, "out/m.xlsx")))["Papers"]
    assert ws.max_row == 3  # the duplicate title is dropped
    assert "method=reward modelling x modality=vision" in result["stdout"]
    assert (
        "[Let's Verify Step by Step](https://arxiv.org/abs/2305.20050)"
        in base64.b64decode(result["files"]["out/m.md"]).decode()
    )


async def test_bibtex_round_trip(run, b64, skill_files):
    refs = [
        {"title": "Let's Verify Step by Step", "authors": ["Hunter Lightman", "Vineet Kosaraju"],
         "year": 2023, "venue": "ICLR 2024", "arxiv": "2305.20050"},
        {"title": "Let's verify step by step", "authors": "Hunter Lightman", "year": 2023},
        {"title": "Attention Is All You Need", "authors": ["Ashish Vaswani"], "year": 2017,
         "venue": "Advances in Neural Information Processing Systems"},
    ]  # fmt: skip
    s = "skills/core/citations/scripts/bib_tools.py"
    code = (
        f"python {s} from-json refs.json out/r.bib\n"
        f"python {s} check out/r.bib || echo CHECK_FAILED\n"
        f"python {s} dedupe out/r.bib out/d.bib\n"
        f"python {s} format out/d.bib --style ieee\n"
    )
    files = {**skill_files("citations"), "refs.json": b64(json.dumps(refs))}
    result = await run(code, files, language="bash")
    bib = out(result, "out/d.bib").decode()
    assert "@inproceedings{lightman2023lets," in bib and bib.count("@") == 2
    assert "same paper: lightman2023lets, lightman2023letsa" in result["stdout"]
    assert "CHECK_FAILED" in result["stdout"]
    assert '[1] H. Lightman and V. Kosaraju, "Let\'s Verify Step by Step,"' in result["stdout"]


DATA = "id,date,amount,group,score\n" + "\n".join(
    f"{i},2025-0{1 + i % 9}-1{i % 9},{100 + i * 3},{'ab'[i % 2]},{(i * 7) % 11 + (i % 2) * 2}"
    for i in range(80)
)


async def test_profile_and_statistics(run, b64, skill_files):
    files = {**skill_files("data-analysis", "statistics"), "in/d.csv": b64(DATA)}
    st = "skills/core/statistics/scripts/stats_tests.py"
    code = (
        "set -e\npython skills/core/data-analysis/scripts/profile.py in/d.csv --out out/p.md\n"
        f"python {st} compare in/d.csv --group group --value score\n"
        f"python {st} proportions --success 120 --total 2400 --success 151 --total 2390\n"
        f"python {st} correlate in/d.csv --x amount --y id\n"
        f"python {st} bootstrap in/d.csv --value score --stat p90 --by group\n"
    )
    result = await run(code, files, language="bash")
    profile = out(result, "out/p.md").decode()
    assert "80 rows x 5 columns" in profile and "dates stored as text" in profile
    assert "Welch's t-test" in result["stdout"] and "Hedges' g" in result["stdout"]
    assert "Relative lift = +26.4%" in result["stdout"]
    assert "Pearson r = 1.000" in result["stdout"]


async def test_ml_baseline(run, b64, skill_files):
    rows = ["x1,x2,colour,label"] + [
        f"{i % 17},{(i * 5) % 13},{'rgb'[i % 3]},{'yes' if (i % 17) > 8 else 'no'}"
        for i in range(240)
    ]
    files = {**skill_files("ml-experiments"), "in/d.csv": b64("\n".join(rows))}
    code = (
        "python skills/core/ml-experiments/scripts/baseline.py in/d.csv --target label "
        "--folds 3 --out out/b"
    )
    result = await run(code, files, timeout_s=180, language="bash")
    report = json.loads(out(result, "out/b/results.json"))
    assert report["task"] == "classification"
    assert report["results"]["random forest"]["roc_auc"]["mean"] > 0.9
    assert report["importances"][0]["feature"] == "x1"


async def test_eval_report(run, b64, skill_files):
    lines = ["system,item_id,score,category"]
    for i in range(60):
        lines += [f"base,{i},{int(i % 3 == 0)},c{i % 2}", f"new,{i},{int(i % 3 != 1)},c{i % 2}"]
    labels = "human,judge\n" + "\n".join(f"{i % 3},{i % 3 if i % 5 else 0}" for i in range(40))
    e = "skills/core/llm-evaluation/scripts/eval_report.py"
    files = {
        **skill_files("llm-evaluation"),
        "r.csv": b64("\n".join(lines)),
        "l.csv": b64(labels),
    }
    code = (
        f"set -e\npython {e} summary r.csv\npython {e} compare r.csv --a base --b new\n"
        f"python {e} agreement l.csv --rater1 human --rater2 judge\n"
    )
    result = await run(code, files, language="bash")
    assert result["exit_code"] == 0, result["stderr"]
    assert "Verdict: new is better" in result["stdout"]
    assert "McNemar exact test" in result["stdout"] and "Cohen's kappa" in result["stdout"]


async def test_free_slots_and_decision_matrix(run, b64, skill_files):
    busy = [
        {"start": "2030-10-21T10:00:00+05:30", "end": "2030-10-21T11:00:00+05:30"},
        {"start": {"dateTime": "2030-10-21T10:30:00+05:30"},
         "end": {"dateTime": "2030-10-21T12:00:00+05:30"}},
        {"start": "2030-10-22", "end": "2030-10-23"},
    ]  # fmt: skip
    spec = {
        "criteria": [
            {"name": "Price", "weight": 50, "better": "lower"},
            {"name": "Quality", "weight": 50, "scale": [1, 5]},
        ],
        "options": [
            {"name": "A", "values": {"Price": 10, "Quality": 3}},
            {"name": "B", "values": {"Price": 20, "Quality": 5}},
            {"name": "C", "values": {"Price": 5, "Quality": 1}, "must_fail": "no support"},
        ],
    }
    files = {
        **skill_files("scheduling", "compare-options"),
        "busy.json": b64(json.dumps(busy)),
        "o.json": b64(json.dumps(spec)),
    }
    code = (
        "set -e\npython skills/core/scheduling/scripts/free_slots.py busy.json --from 2030-10-21 "
        "--to 2030-10-22 --tz Asia/Kolkata --hours 09:30-18:30 --minutes 45 --also Europe/Berlin\n"
        "python skills/core/compare-options/scripts/decision_matrix.py o.json --xlsx out/c.xlsx\n"
    )
    result = await run(code, files, language="bash")
    assert (
        "- Mon 21 Oct 12:00-18:30 (390 min free) | Europe/Berlin: 08:30-15:00" in result["stdout"]
    )
    assert "Tue 22 Oct 09:30-18:30" in result["stdout"]  # all-day events don't block by default
    assert "| 1 | A | 75 |" in result["stdout"] and "A stays first" in result["stdout"]
    assert "C (fails: no support)" in result["stdout"]
    assert (
        load_workbook(io.BytesIO(out(result, "out/c.xlsx")))["Comparison"]["A1"].value == "Option"
    )


async def test_images_and_conversion(run, b64, skill_files):
    buf = io.BytesIO()
    Image.new("RGB", (1200, 800), "teal").save(buf, "JPEG")
    doc = Document()
    doc.add_heading("Report", 1)
    doc.add_paragraph("Body text")
    dbuf = io.BytesIO()
    doc.save(dbuf)
    files = {
        **skill_files("images", "file-conversion"),
        "in/p.jpg": b64(buf.getvalue()),
        "in/r.docx": b64(dbuf.getvalue()),
        "in/t.csv": b64("a;b\n1;x\n2;y|z\n"),
    }
    i, c = (
        "skills/core/images/scripts/image_tools.py",
        "skills/core/file-conversion/scripts/convert.py",
    )
    code = (
        f"set -e\npython {i} resize in/p.jpg out/small.jpg --max 600\n"
        f"python {i} crop in/p.jpg out/sq.png --aspect 1:1\n"
        f"python {i} convert in/p.jpg out/sq.png out/both.pdf\n"
        f"python {i} watermark in/p.jpg out/w.jpg 'Draft'\n"
        f"python {i} strip in/p.jpg out/clean.jpg\n"
        f"python {c} in/t.csv out/t.xlsx\npython {c} in/t.csv out/t.md\n"
        f"python {c} in/r.docx out/r.md\n"
    )
    result = await run(code, files, language="bash")
    assert Image.open(io.BytesIO(out(result, "out/small.jpg"))).size == (600, 400)
    assert Image.open(io.BytesIO(out(result, "out/sq.png"))).size == (800, 800)
    assert len(PdfReader(io.BytesIO(out(result, "out/both.pdf"))).pages) == 2
    assert load_workbook(io.BytesIO(out(result, "out/t.xlsx"))).active["B3"].value == "y|z"
    assert "| 2 | y\\|z |" in out(result, "out/t.md").decode()
    assert out(result, "out/r.md").decode().startswith("# Report\n\nBody text")
