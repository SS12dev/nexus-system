"""Every script shipped in library/core runs in the sandbox as SKILL.md describes."""

import base64
import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

REPORT_MD = """# EV battery suppliers

The market is **concentrated**: see [IEA](https://www.iea.org).

| Supplier | Capacity (GWh) |
| --- | --- |
| CATL | 400 |
| BYD | 250 |

- China leads
  - CATL first
- Korea second

1. First
2. Second

![chart](chart.png)
"""

CHART = {
    "type": "bar",
    "title": "Capacity",
    "x": ["CATL", "BYD"],
    "series": [{"name": "GWh", "values": [400, 250]}],
    "source": "test",
}


def out(result, path) -> bytes:
    assert result["exit_code"] == 0, result["stderr"]
    return base64.b64decode(result["files"][path])


def _has_pango() -> bool:
    try:
        import weasyprint  # noqa: F401
    except OSError:
        return False
    return True


# WeasyPrint needs Pango, which the sandbox image installs; on a Mac: brew install pango.
@pytest.mark.skipif(not _has_pango(), reason="WeasyPrint needs the Pango system library")
async def test_chart_then_pdf_report(run, b64, skill_files):
    files = {
        **skill_files("pdf", "charts"),
        "work/report.md": b64(REPORT_MD),
        "work/chart.json": b64(json.dumps(CHART)),
    }
    code = (
        "python skills/core/charts/scripts/chart.py work/chart.json work/chart.png\n"
        "python skills/core/pdf/scripts/md_to_pdf.py work/report.md out/report.pdf --title EV\n"
        "python skills/core/pdf/scripts/pdf_tools.py merge out/both.pdf out/report.pdf "
        "out/report.pdf\n"
        "python skills/core/pdf/scripts/pdf_tools.py watermark out/both.pdf out/draft.pdf DRAFT\n"
        "python skills/core/pdf/scripts/pdf_tools.py tables out/report.pdf out/tables.xlsx\n"
    )
    result = await run(code, files, language="bash")
    pdf = PdfReader(io.BytesIO(out(result, "out/report.pdf")))
    text = pdf.pages[0].extract_text()
    assert "EV battery suppliers" in text and "CATL" in text
    assert pdf.metadata.title == "EV"
    assert len(PdfReader(io.BytesIO(out(result, "out/draft.pdf"))).pages) == 2 * len(pdf.pages)
    wb = load_workbook(io.BytesIO(out(result, "out/tables.xlsx")))
    assert wb.worksheets[0]["A2"].value == "CATL"


async def test_fill_pdf_form(run, b64, skill_files):
    make_form = """
from reportlab.pdfgen import canvas
c = canvas.Canvas("in/form.pdf")
c.acroForm.textfield(name="company", x=100, y=700, width=200, height=20)
c.acroForm.checkbox(name="agree", x=100, y=650)
c.showPage()
c.save()
"""
    files = {
        **skill_files("pdf"),
        "values.json": b64(json.dumps({"company": "Acme", "agree": True})),
    }
    code = (
        f"python - <<'EOF'\n{make_form}\nEOF\n"
        "python skills/core/pdf/scripts/fill_form.py fields in/form.pdf\n"
        "python skills/core/pdf/scripts/fill_form.py fill in/form.pdf out/filled.pdf values.json\n"
    )
    result = await run(code, files, language="bash")
    assert "company" in result["stdout"], result["stderr"]
    fields = PdfReader(io.BytesIO(out(result, "out/filled.pdf"))).get_fields()
    assert fields["company"]["/V"] == "Acme"
    assert fields["agree"]["/V"] != "/Off"


async def test_docx_create_replace_append(run, b64, skill_files):
    files = {
        **skill_files("docx"),
        "work/report.md": b64(REPORT_MD.replace("![chart](chart.png)", "Price: 10 USD")),
        "work/map.json": b64(json.dumps({"10 USD": "12 USD"})),
        "work/more.md": b64("## Appendix\n\nExtra."),
    }
    code = (
        "set -e\n"
        "python skills/core/docx/scripts/md_to_docx.py work/report.md out/r.docx --title T --toc\n"
        "python skills/core/docx/scripts/docx_tools.py replace out/r.docx out/r2.docx "
        "work/map.json\n"
        "python skills/core/docx/scripts/docx_tools.py append out/r2.docx out/r3.docx "
        "work/more.md\n"
        "python skills/core/docx/scripts/docx_tools.py inspect out/r3.docx\n"
    )
    result = await run(code, files, language="bash")
    doc = Document(io.BytesIO(out(result, "out/r3.docx")))
    styles = [p.style.name for p in doc.paragraphs]
    assert "Heading 1" in styles and "List Bullet 2" in styles and "List Number" in styles
    assert any("12 USD" in p.text for p in doc.paragraphs)
    assert doc.paragraphs[-1].text == "Extra."
    assert doc.tables[0].cell(1, 0).text == "CATL"


async def test_pptx_build_edit(run, b64, skill_files):
    deck = {
        "title": "Q3",
        "slides": [
            {"layout": "title", "title": "Q3 update", "subtitle": "Oct"},
            {"layout": "bullets", "title": "Sales up 18%", "bullets": ["a", ["b"]], "notes": "n"},
            {"layout": "image", "title": "Chart", "image": "chart.png", "caption": "src"},
            {"layout": "table", "title": "Top", "rows": [["A", "B"], ["Acme", "1"]]},
            {"layout": "two_column", "title": "Risks", "left": ["x"], "right": ["y"]},
            {"layout": "section", "title": "End"},
        ],
    }
    files = {
        **skill_files("pptx", "charts"),
        "work/deck.json": b64(json.dumps(deck)),
        "work/chart.json": b64(json.dumps(CHART)),
        "work/map.json": b64(json.dumps({"18%": "20%", "Acme": "Globex"})),
    }
    code = (
        "set -e\n"
        "python skills/core/charts/scripts/chart.py work/chart.json work/chart.png\n"
        "python skills/core/pptx/scripts/build_deck.py work/deck.json out/deck.pptx\n"
        "python skills/core/pptx/scripts/pptx_tools.py replace out/deck.pptx out/d2.pptx "
        "work/map.json\n"
        "python skills/core/pptx/scripts/pptx_tools.py delete out/d2.pptx out/d3.pptx 6\n"
        "python skills/core/pptx/scripts/build_deck.py work/deck.json out/t.pptx "
        "--template out/deck.pptx\n"
    )
    result = await run(code, files, language="bash")
    prs = Presentation(io.BytesIO(out(result, "out/d3.pptx")))
    assert len(prs.slides) == 5
    assert prs.slides[1].shapes.title.text == "Sales up 20%"
    assert prs.slides[1].notes_slide.notes_text_frame.text == "n"
    assert len(Presentation(io.BytesIO(out(result, "out/t.pptx"))).slides) == 6


async def test_xlsx_build_inspect(run, b64, skill_files):
    book = {
        "sheets": [
            {
                "name": "Suppliers",
                "columns": [
                    {"header": "Supplier"},
                    {"header": "GWh", "format": "#,##0"},
                    {"header": "Since"},
                ],
                "rows": [["CATL", 400, "2011-12-16"], ["BYD", 250, "1995-02-10"]],
                "totals": {"GWh": "sum"},
                "chart": {"type": "col", "categories": "Supplier", "values": ["GWh"]},
            }
        ]
    }
    files = {**skill_files("xlsx"), "book.json": b64(json.dumps(book))}
    code = (
        "set -e\npython skills/core/xlsx/scripts/build_workbook.py book.json out/b.xlsx\n"
        "python skills/core/xlsx/scripts/xlsx_tools.py inspect out/b.xlsx\n"
    )
    result = await run(code, files, language="bash")
    ws = load_workbook(io.BytesIO(out(result, "out/b.xlsx")))["Suppliers"]
    assert ws["B4"].value == "=SUM(B2:B3)"
    assert ws.freeze_panes == "A2" and len(ws._charts) == 1
    assert "=SUM(B2:B3)" in result["stdout"]


async def test_every_chart_type_renders(run, b64, skill_files):
    files = skill_files("charts")
    for kind in ("bar", "barh", "stacked_bar", "line", "area", "pie", "donut", "scatter"):
        files[f"{kind}.json"] = b64(json.dumps({**CHART, "type": kind}))
    code = "\n".join(
        f"python skills/core/charts/scripts/chart.py {k}.json out/{k}.svg"
        for k in ("bar", "barh", "stacked_bar", "line", "area", "pie", "donut", "scatter")
    )
    result = await run("set -e\n" + code, files, language="bash")
    assert result["exit_code"] == 0, result["stderr"]
    assert len(result["files"]) == 8
